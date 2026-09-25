"""Independent check of convert_workhere_refs.py: original (Harvard) vs converted (Vancouver) working copy.

1. every paragraph without citations and every table cell is byte-identical;
2. in paragraphs with citations, the text is identical once each citation group is replaced by a placeholder;
3. citation groups, in document order, point to the same papers (first author + year);
4. the numbered list is 1..N, every number is cited, numbers follow first citation;
5. each numbered entry keeps the original entry's title, DOI and italic runs;
6. no author-date citation is left.
Usage: python verify_workhere_refs.py ORIGINAL.docx CONVERTED.docx
"""
import re
import sys

import docx
from docx.oxml.ns import qn
from docx.text.run import Run

orig, conv = docx.Document(sys.argv[1]), docx.Document(sys.argv[2])
fails = []


def check(label, ok, detail=""):
    print(f"{'PASS' if ok else 'FAIL'}  {label}" + (f"  |  {detail}" if detail else ""))
    if not ok:
        fails.append(label)


def runs(p):
    return [Run(r, p) for r in p._p.iter(qn("w:r"))]


def text(p):
    return "".join(r.text for r in runs(p))


NAME = r"[A-Z][A-Za-zÀ-ɏ'’\-]+(?: [A-Z][A-Za-zÀ-ɏ'’\-]+)*"
AD_ITEM = re.compile(rf"({NAME})(?: and {NAME}| et al\.)?, ((?:19|20)\d\d)[a-z]?(?=;|\)|$)")
AD_GROUP = re.compile(rf"\([^()]*?, (?:19|20)\d\d[a-z]?(?=[;)])[^()]*\)|{NAME}(?: and {NAME}| et al\.)? \((?:19|20)\d\d[a-z]?\)")
NUM_GROUP = re.compile(r"(?:\([^()]*\) )?\[(\d+(?:\s*[,–]\s*\d+)*)\]|" + NAME + r"(?: and " + NAME + r"| et al\.)? \[(\d+)\]")


def split_refs(d):
    P = d.paragraphs
    ri = [i for i, p in enumerate(P) if p.text.strip() == "References"][0]
    return P[:ri], P[ri + 1:]


ob, orest = split_refs(orig)
cb, crest = split_refs(conv)
check("same number of paragraphs before References", len(ob) == len(cb), f"{len(ob)} vs {len(cb)}")

# 1-2 body text
intro = [i for i, p in enumerate(ob) if re.fullmatch(r"(?:1\. )?Introduction", p.text.strip())][0]
ogroups, changed_other = [], []
# front matter: citations removed (names kept for narrative ones), nothing else changed, no bracket numbers
front_bad = []
for i, (po, pc) in enumerate(zip(ob[:intro], cb[:intro])):
    to, tc = text(po), text(pc)
    exp = re.sub(rf"({NAME}(?: and {NAME}| et al\.)?) \((?:19|20)\d\d[a-z]?\)", r"\1", to)
    exp = re.sub(rf" ?\((?:{NAME}(?: and {NAME}| et al\.)?, (?:19|20)\d\d[a-z]?(?:; )?)+\)", "", exp)
    if exp != tc or re.search(r"\[\d", tc):
        front_bad.append(i)
check("front matter: citations removed, nothing else changed", not front_bad, f"paragraphs: {front_bad or 'none'}")
for i, (po, pc) in enumerate(zip(ob[intro:], cb[intro:]), intro):
    to, tc = text(po), text(pc)
    go = list(AD_GROUP.finditer(to))
    go = [m for m in go if AD_ITEM.search(m.group(0).strip("()")) or re.search(r"\((?:19|20)\d\d\)$", m.group(0))]
    if not go:
        if to != tc:
            changed_other.append(i)
        continue
    for m in go:
        g = m.group(0)
        if g.startswith("("):
            ks = [(a, y) for a, y in AD_ITEM.findall(g[1:-1])]
        else:
            a = re.match(NAME, g).group(0)
            ks = [(a, re.search(r"\((\d{4})", g).group(1))]
        ogroups.append(sorted(ks))
    def _ph(m):  # original group -> placeholder, keeping non-citation items as the converter does
        g = m.group(0)
        if not g.startswith("("):
            return g[:g.rindex(" (")] + " §"
        keep = [x.strip() for x in g[1:-1].split(";") if not AD_ITEM.fullmatch(x.strip())]
        return (f"({'; '.join(keep)}) " if keep else "") + "§"
    conv_ph = re.sub(r"\[\d+(?:\s*[,–]\s*\d+)*\]", "§", tc)
    if AD_GROUP.sub(_ph, to) != conv_ph:
        changed_other.append(i)
check("no text outside citations changed (body)", not changed_other, f"paragraphs: {changed_other or 'none'}")

ot = [c.text for t in orig.tables for r in t.rows for c in r.cells]
ct = [c.text for t in conv.tables for r in t.rows for c in r.cells]
check("tables unchanged", ot == ct)

# reference lists
def entries(rest, numbered):
    out = []
    for p in rest:
        t = text(p).strip()
        if not t:
            continue
        if numbered and not re.match(r"^\d+\. ", t):
            break
        if not numbered and (not re.match(r"^.+?, \d{4}[a-z]?\. ", t) or t.startswith("Table ")):
            break
        out.append(p)
    return out


oe, ce = entries(orest, False), entries(crest, True)
nums = [int(re.match(r"(\d+)\. ", text(p).strip()).group(1)) for p in ce]
check("numbered list is 1..N", nums == list(range(1, len(nums) + 1)), f"{len(nums)} entries")
check("same number of entries as the original", len(oe) == len(ce), f"{len(oe)} vs {len(ce)}")

okey = {}
for p in oe:
    t = text(p).strip()
    y = re.search(r", (\d{4})[a-z]?\. ", t)
    title_start = y.end()
    doi = re.search(r"https?://\S+$", t).group(0)
    ital = [r.text for r in runs(p) if r.italic and r.text.strip()]
    okey[(t.split(", ")[0], y.group(1))] = (t, doi, ital[:-1] if "doi.org" in doi else ital, title_start)
vkey, bad = {}, []
for p in ce:
    t = text(p).strip()
    n = int(t.split(".")[0])
    fa = re.match(r"\d+\. (.+?) [A-Z]+[,.]", t).group(1)
    yr = re.findall(r"[ (]((?:19|20)\d\d)[;)]", t)[-1]
    vkey[n] = (fa, yr)
    o = okey.get((fa, yr))
    if not o:
        bad.append((n, "no original"))
        continue
    ot_, odoi, oital, ts = o
    if odoi not in t:
        bad.append((n, "doi"))
    ital = [r.text for r in runs(p) if r.italic and r.text.strip()]
    if ital != oital:
        bad.append((n, f"italics {oital} -> {ital}"))
check("every entry maps to an original entry with the same DOI and italic title runs", not bad, f"{bad or 'none'}")
titles_ok = []
for p in ce:
    t = text(p).strip()
    fa, yr = vkey[int(t.split(".")[0])]
    ot_ = okey[(fa, yr)][0]
    ts = okey[(fa, yr)][3]
    # the original title is everything from the year to the italic journal; require it verbatim in the new entry
    otitle = re.sub(r" [^.?!]+? \d+(?:\(\d+\))?, [A-Za-z]*\d+(?:–[A-Za-z]*\d+)?\.(?: https://doi\.org/\S+)?$", "", ot_[ts:])
    otitle = re.sub(r" \S+\. https?://\S+$", "", otitle)
    if otitle not in t:
        titles_ok.append(int(t.split(".")[0]))
check("every entry keeps its original title text", not titles_ok, f"{titles_ok or 'none'}")

# 3-4 citation groups
cgroups, flat = [], []
for p in cb:
    for m in NUM_GROUP.finditer(text(p)):
        g = m.group(1) or m.group(2)
        ns = []
        for part in g.split(","):
            a, _, b = part.strip().partition("–")
            ns += list(range(int(a), int(b or a) + 1))
        flat += ns
        cgroups.append(sorted(vkey[n] for n in ns))
check("citation groups point to the same papers in the same order", ogroups == cgroups,
      f"{len(cgroups)} groups vs {len(ogroups)} original")
first = list(dict.fromkeys(flat))
check("every entry cited; numbers follow order of first citation", first == nums, f"{len(first)} cited")
left = [text(p) for p in cb if re.search(r"(?:et al\.|[A-Z][a-z]+), (?:19|20)\d\d[a-z]?[;)]", text(p))]
check("no author-date citation left", not left)

print(f"\n{'ALL CHECKS PASS' if not fails else str(len(fails)) + ' FAILED: ' + ', '.join(fails)}")
sys.exit(1 if fails else 0)
