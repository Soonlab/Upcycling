#!/usr/bin/env python
"""Build the EndNote-ready manuscript: 01_Manuscript_EndNote.docx

../01_Manuscript.md stays the master and is not modified. This script
  1. rewrites every in-text citation as an EndNote temporary citation {Family, Year},
     using the family name / year held in the library (citekeys.json), so that
     EndNote's "Update Citations and Bibliography" links them to the imported records;
  2. marks any citation whose reference is still `pending` in ref_decisions.tsv as
     **[UNVERIFIED-REF: ...]** (yellow highlight) instead of linking it; a citation that
     matches no library record at all stops the build;
  3. drops the typed reference list and moves the tables in front of the References
     heading, so the bibliography EndNote appends at the end of the file lands under it.

Non-citation text that shared a parenthesis with a citation, e.g. "(r220; Chaumeil et
al., 2022)", is split into "(r220) {Chaumeil, 2022}" - prefix/suffix syntax inside
temporary citations is avoided on purpose.
"""
import json, re, subprocess, sys, unicodedata
from pathlib import Path

HERE = Path(__file__).resolve().parent
MD = HERE.parent / "01_Manuscript.md"
OUT_MD = HERE / "01_Manuscript_EndNote.md"
OUT_DOCX = HERE / "01_Manuscript_EndNote.docx"
PANDOC = "/home/soon/miniconda3/bin/pandoc"

NAME = r"[A-Z][A-Za-zÀ-ÿ'’\-]+"
SEG = re.compile(rf"^\s*({NAME})(?: et al\.| and {NAME})?,\s*((?:19|20)\d{{2}})[a-z]?\s*$")


def fold(s):
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()


def resolve(surname, year, keys):
    hits = [n for n, k in keys.items() if fold(k["family"]) == fold(surname) and k["year"] == year]
    return (hits[0], "") if len(hits) == 1 else (None, f"cannot resolve {surname} {year}")


def main():
    text = MD.read_text(encoding="utf-8")
    keys = {int(k): v for k, v in json.loads((HERE / "citekeys.json").read_text()).items()}
    assert len(keys) > 0, "citekeys.json is empty - run 02_build_library.py"
    families = {fold(k["family"]) for k in keys.values()}

    head, rest = text.split("\n## References", 1)
    _refs, tail = rest.split("\n---", 1)  # tail = tables
    log, used, unresolved = [], {}, []

    def convert(m):
        segs = m.group(1).split(";")
        if not any(SEG.match(s) for s in segs):
            return m.group(0)
        parts, run = [], []  # run = consecutive verified citations

        def flush():
            if run:
                parts.append("\\{" + "; ".join(run) + "\\}")
                run.clear()

        for s in segs:
            c = SEG.match(s)
            if not c:
                flush()
                parts.append("(" + s.strip() + ")")
                continue
            n, warn = resolve(c.group(1), c.group(2), keys)
            if n is None:
                unresolved.append(s.strip())
                flush()
                parts.append(f"**[UNVERIFIED-REF: {s.strip()}]**")
                continue
            used[n] = used.get(n, 0) + 1
            k = keys[n]
            if warn or (k["status"] != "pending" and k["year"] != c.group(2)):
                log.append(f"R{n}: '{s.strip()}' -> {warn or 'library year ' + k['year']}")
            if k["status"] == "pending":
                flush()
                parts.append(f"**[UNVERIFIED-REF: {s.strip()}]**")
            else:
                run.append(f"{k['family']}, {k['year']}")
        flush()
        return " ".join(parts)

    body = re.sub(r"\(([^()]*)\)", convert, head)

    # every author-year string in the body must have been converted
    stripped = re.sub(r"\\\{[^{}]*\\\}|\*\*\[UNVERIFIED-REF:[^\]]*\]\*\*", "", body)
    left = [m.group(0) for m in re.finditer(rf"({NAME})(?: et al\.| and {NAME})?,? \(?(?:19|20)\d{{2}}", stripped)
            if fold(m.group(1)) in families]
    assert not unresolved, f"unresolved citations: {unresolved}"
    assert not left, f"author-year strings left unconverted: {left}"

    out = (body.rstrip() + "\n\n---\n" + tail.rstrip() + "\n\n---\n\n## References\n\n")
    OUT_MD.write_text(out, encoding="utf-8")
    subprocess.run([PANDOC, str(OUT_MD), "-o", str(OUT_DOCX), "--from=markdown+pipe_tables+tex_math_dollars"], check=True)

    # highlight the unverified markers
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    doc = Document(str(OUT_DOCX))
    n_marked = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if "UNVERIFIED-REF" in r.text:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                n_marked += 1
    doc.save(str(OUT_DOCX))
    full = "\n".join(p.text for p in doc.paragraphs)
    n_temp = len(re.findall(r"\{[^{}]+\}", full))
    n_cites = sum(len(x.split(";")) for x in re.findall(r"\{([^{}]+)\}", full))

    cited = set(used)
    not_cited = sorted(n for n, k in keys.items() if n not in cited)
    pend_cited = sorted(n for n in cited if keys[n]["status"] == "pending")
    rep = [f"temporary citation groups in docx : {n_temp}  (individual citations {n_cites})",
           f"unverified markers highlighted    : {n_marked}  (refs {pend_cited})",
           f"library records cited             : {len([n for n in cited if keys[n]['status'] != 'pending'])}",
           f"listed but never cited in text    : {not_cited}",
           "year / key notes:"] + ["   " + x for x in sorted(set(log))]
    (HERE / "endnote_docx_report.txt").write_text("\n".join(rep) + "\n", encoding="utf-8")
    print("\n".join(rep))
    assert n_temp > 0


if __name__ == "__main__":
    main()
